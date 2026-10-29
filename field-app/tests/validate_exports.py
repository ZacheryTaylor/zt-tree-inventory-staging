"""Validate sample exports: PNEZD is valid 5-column CSV with numeric N/E/Z, tree log parses, XLSX opens,
and an independent library (pyproj/PROJ) reproduces the northing/easting from the logged lat/long."""
import csv, json, os, sys
from pyproj import Transformer
import openpyxl
D = os.path.join(os.path.dirname(__file__), '..', 'sample')
res = {}
import glob, re
pn_files = glob.glob(os.path.join(D, 'sample_PNEZD_v*.csv'))
assert len(pn_files) == 1, pn_files  # version in file name
PN = pn_files[0]
rows = list(csv.reader(open(PN, newline='')))
notes = open(PN[:-4] + '_NOTES.txt').read()
for must in ['ZT Tree Inventory v1.0.1', 'ISA Certified Arborist', 'not a survey', 'Not a survey', 'not survey grade', 'Tool, not engineering or surveying advice', 'Zach Taylor, EI (Engineer Intern), ZT, LLC', 'EPSG:2238', 'US survey feet', 'gps.gov']:
    assert must in notes, must
assert len(rows) == 5, rows
for r in rows:
    assert len(r) == 5, r
    int(r[0]); float(r[1]); float(r[2]); float(r[3])
    assert ',' not in r[4] and r[4].startswith('T')
assert len({r[0] for r in rows}) == 5
res['pnezd'] = {'file': os.path.basename(PN), 'rows': len(rows), 'cols': 5, 'header_row': False, 'first': rows[0], 'notes_file_ok': True}
lines = list(csv.reader(open(os.path.join(D, 'sample_tree_log.csv'), newline='')))
hi = next(i for i, r in enumerate(lines) if r and r[0] == 'Tag #')
hdr, data = lines[hi], lines[hi + 1:]
assert len(data) == 5 and all(len(r) == len(hdr) for r in data)
t = Transformer.from_crs('EPSG:4326', 'EPSG:2238', always_xy=True)
maxd = 0
pn = {r[0]: r for r in rows}
for r in data:
    d = dict(zip(hdr, r)); e, n = t.transform(float(d['Longitude']), float(d['Latitude']))
    maxd = max(maxd, abs(e - float(d['Easting (ft)'])), abs(n - float(d['Northing (ft)'])))
    p = pn[d['Point #']]; assert abs(float(p[1]) - float(d['Northing (ft)'])) < 1e-6 and d['Tag #'] in p[4]
man = [dict(zip(hdr, r)) for r in data if r[hdr.index('Position source')].startswith('Manual')]
assert len(man) == 1 and man[0]['GPS accuracy (ft, reported)'] == '' and man[0]['Spread (ft)'] == '', man
gps = [dict(zip(hdr, r)) for r in data if r[hdr.index('Position source')].startswith('Phone GPS')]
assert all(float(g['GPS accuracy (ft, reported)']) > 0 for g in gps)
for must in ['Tool, not engineering or surveying advice', 'Zach Taylor, EI (Engineer Intern), ZT, LLC', 'Not a survey', 'gps.gov']:
    assert any(must in ' '.join(l) for l in lines[:hi]), must
res['tree_log'] = {'manual_rows_accuracy_blank': True, 'rows': len(data), 'cols': len(hdr), 'keyed_by': hdr[0], 'pyproj_max_diff_ft': round(maxd, 4), 'quoted_comma_note_ok': dict(zip(hdr, data[4]))['Field notes'],
    'suggested_cols': [h for h in hdr if h.startswith('Suggest')]}
assert 'Proposed action' not in hdr, 'tree log must not carry Proposed action'
assert hdr[-4:] == ['Suggested protected (code threshold)', 'Suggested heritage (code threshold)', 'Suggestion basis', 'Accuracy flag'], hdr[-4:]
assert dict(zip(hdr, data[0]))['Suggested protected (code threshold)'] == 'Y' and dict(zip(hdr, data[1]))['Suggested protected (code threshold)'] == 'N'  # ESC: live oak 32 in Y; longleaf pine exempt
assert any('field judgement' in ' '.join(l) for l in lines[:hi])
assert maxd < 0.01
wb = openpyxl.load_workbook(os.path.join(D, 'sample_tree_log.xlsx'))
res['xlsx'] = {s: wb[s].max_row for s in wb.sheetnames}
assert wb.sheetnames == ['Tree Log', 'Calculator paste', 'PNEZD', 'About', 'Accuracy checks'], wb.sheetnames
ac = list(wb['Accuracy checks'].iter_rows(values_only=True))
assert ac[1][7] == 'Horizontal offset (ft)' and abs(float(ac[2][7]) - 12.02) < 0.01 and ac[2][10] == 'PASS', ac[2]
res['accuracy_checks_sheet'] = {'offset_ft': ac[2][7], 'result': ac[2][10]}
cp = list(wb['Calculator paste'].iter_rows(values_only=True))
calc_hdr = ["Tag #", "Species", "DBH (in)", "Dripline radius (ft)", "Condition", "Proposed action", "In jurisdictional wetland? (PNS)",
  "Planted/preserved to meet prior approval? (ESC)", "ESC sand live oak multi-stem qualifies? (>=5 stems or >=3 stems >=3 in)",
  "SR champion tree?", "Canopy to be pruned (%)", "Dripline / root zone disturbed (%)", "Field notes"]  # tree-mitigation-calculator p4_inventory.COLS[:13]
assert list(cp[1]) == calc_hdr, cp[1]
assert cp[2][6:12] == ('N', 'N', 'N', 'N', 10, 20), cp[2]
assert all(r[5] in (None, '') for r in cp[2:]), 'Proposed action column must be blank'
assert 'Proposed action' in cp[0][0] and 'column-for-column' in cp[0][0], cp[0]
assert not any('Proposed action' in str(c) for row in wb['Tree Log'].iter_rows(values_only=True) for c in row if c), 'Tree Log must not carry Proposed action'
res['calc_paste_action_blank_with_note'] = True
res['calc_paste'] = {'header_matches_calculator': True, 'row1_G_L': list(cp[2][6:12])}
disc = next(l for l in notes.splitlines() if l.startswith('Tool, not engineering'))  # = core.js DISCLAIMER
plan = open(os.path.join(os.path.dirname(__file__), '..', 'PLAN.md')).read()
assert disc in plan, 'PLAN.md disclaimer differs from app DISCLAIMER'
res['plan_disclaimer_matches_app'] = True
bk = json.load(open(os.path.join(D, 'sample_job_backup.json')))
assert bk['app'] == 'zt-tree-inventory' and bk['version'] == '1.0.1' and not any('action' in t for t in bk['trees']) and len(bk['trees']) == 5 and bk['job']['jurisdiction'] == 'ESC'
res['job_backup'] = {'trees': len(bk['trees']), 'checks': len(bk['checks']), 'parcel_cached': bool(bk['parcelGeo'])}
alltext = notes + open(os.path.join(D, 'sample_tree_log.csv')).read()
assert 'Tree Survey' not in alltext and 'tree survey' not in alltext.lower(), 'banned term in outputs'
json.dump(res, open(os.path.join(os.path.dirname(__file__), 'validate_exports_results.json'), 'w'), indent=2)
print(json.dumps(res, indent=2)); print('ALL EXPORT CHECKS PASS')
